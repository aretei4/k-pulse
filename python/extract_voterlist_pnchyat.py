"""
Extract an Odisha electoral-roll PDF (Crystal Reports layout, as in
P030028.pdf) into the K-Pulse voter upload sheet:

    house_no | name | relation | Relation_Name | Age | Gender | Assembly Part | Epic_no | Ward_No

Why not just read the PDF text? The Odia names are set in an embedded
OROT-Mukta subset whose text layer maps every glyph back to a bare consonant:
vowel signs and conjuncts are lost ("ସାହୁ" extracts as "ସନହହ"). So this reads
the glyphs themselves. Each glyph is identified by a hash of its outline,
which stays the same across PDFs even though the glyph numbering changes, and
is looked up in odia_glyphs.json. The glyphs are then put back into logical
Unicode order: the pre-base ୋ/ୌ/ୈ, the reph and the conjunct parts.

Usage:
    pip install pymupdf fonttools openpyxl
    python extract_voterlist.py P030028.pdf                    -> P030028.xlsx
    python extract_voterlist.py P030028.pdf -o out.xlsx --ward 4 --pages 3-10
    python extract_voterlist.py *.pdf                           (one sheet per PDF)

Glyphs missing from the table fall back to the PDF's own (lossy) letter and are
reported. Run with --dump-unknown DIR to render them as a PNG so they can be
added to odia_glyphs.json.
"""
import argparse
import glob
import hashlib
import io
import json
import re
import sys
from collections import Counter, defaultdict
from pathlib import Path

import pymupdf
from fontTools.ttLib import TTFont
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

HERE = Path(__file__).resolve().parent
HEADERS = ["house_no", "name", "relation", "Relation_Name", "Age", "Gender", "Assembly Part", "Epic_no", "Ward_No"]
WIDTHS = [16.5, 24, 9, 24, 8, 8, 16, 20, 9]

# Column x-ranges (PDF points) of the roll's table, and a row's vertical band
# relative to its serial number. Taken from the Kendrapara rolls.
COLS = {
    "serial": (25, 62), "section": (70, 110), "house": (120, 160),
    "name": (160, 300), "reltype": (270, 300), "relname": (305, 410),
    "gender": (410, 440), "age": (440, 470), "epic": (470, 575),
}
BAND = (-4, 14)

# ---------------------------------------------------------------- glyph → text

CONS = set("କଖଗଘଙଚଛଜଝଞଟଠଡଢଣତଥଦଧନପଫବଭମଯରଲଳଵଶଷସହୟୱଡ଼ଢ଼")
VOWELS = set("ଅଆଇଈଉଊଋଏଐଓଔ")
HALANT, E_SIGN, REPH = "୍", "େ", "ର୍"
# େ followed by a length mark is one two-part vowel sign.
TWO_PART = {"ା": "ୋ", "ୗ": "ୌ", "ୖ": "ୈ"}


def is_sign(ch):
    return "଼" <= ch <= "ୣ" or ch in "ଁଂଃ"


def to_logical(glyphs):
    """Glyph labels in visual order -> logical Unicode.

    Each item is (text, is_reph). The shaper draws the େ before its
    consonant, the reph after the base, and may draw ି before a subjoined
    consonant, so everything is regrouped per syllable cluster.
    """
    out, cl, pending_e = [], None, False

    def new():
        return {"cons": "", "marks": [], "reph": False, "e": False}

    def flush():
        nonlocal cl
        if cl is None:
            return
        text = (REPH if cl["reph"] else "") + cl["cons"]
        marks = cl["marks"]
        if cl["e"]:
            if marks and marks[0] in TWO_PART:
                text += TWO_PART[marks[0]]
                marks = marks[1:]
            else:
                text += E_SIGN
        out.append(text + "".join(marks))
        cl = None

    for text, reph in glyphs:
        if reph:
            if cl is None:
                cl = new()
            cl["reph"] = True
            if text:
                cl["marks"].append(text)
            continue
        if not text:
            continue
        if text == E_SIGN:
            pending_e = True
            continue
        first = text[0]
        if first in CONS:
            if cl is not None and cl["cons"].endswith(HALANT) and not cl["marks"]:
                cl["cons"] += text
            else:
                flush()
                cl = new()
                cl["cons"] = text
            if pending_e:
                cl["e"], pending_e = True, False
        elif first == HALANT:                      # subjoined consonant or bare halant
            if cl is None:
                cl = new()
            cl["cons"] += text
        elif is_sign(first):
            if cl is None:
                out.append(text)
            else:
                cl["marks"].append(text)
        else:
            flush()
            if first in VOWELS:
                cl = new()
                cl["cons"] = text
            else:
                out.append(text)
    flush()
    if pending_e:
        out.append(E_SIGN)
    return re.sub(r"\s+", " ", "".join(out)).strip()


def outline_hash(glyf, name):
    coords, ends, flags = glyf[name].getCoordinates(glyf)
    return hashlib.sha1(repr((list(coords), list(ends), [f & 1 for f in flags])).encode()).hexdigest()[:12]


class Decoder:
    def __init__(self, table_path):
        self.table = json.loads(Path(table_path).read_text(encoding="utf-8"))
        self.unknown = Counter()        # hash -> count
        self.unknown_seen = {}          # hash -> (pdf, font xref, glyph name, example)

    def fonts(self, doc, pdf):
        """font name -> {glyph id -> outline hash}, for every Odia font in the PDF."""
        maps = {}
        for page in doc:
            for xref, _ext, _type, base, *_ in page.get_fonts(full=True):
                name = base.split("+")[-1]
                if "OROT" not in name or name in maps:
                    continue
                font = TTFont(io.BytesIO(doc.extract_font(xref)[3]))
                glyf = font["glyf"]
                order = font.getGlyphOrder()
                maps[name] = {gid: (outline_hash(glyf, gn), (pdf, xref, gn)) for gid, gn in enumerate(order)}
        return maps

    def text(self, chars):
        glyphs = []
        for c in chars:
            h, origin = c["glyph"]
            label = self.table.get(h)
            if label is None:
                self.unknown[h] += 1
                self.unknown_seen.setdefault(h, origin)
                label = c["ch"]                    # the PDF's own guess: right consonant, no signs
            if label.startswith("@"):
                glyphs.append((label[1:], True))
            else:
                glyphs.append((label, False))
        return to_logical(glyphs)


# ---------------------------------------------------------------- page → rows

def page_chars(page, font_maps):
    out = []
    for span in page.get_texttrace():
        gmap = font_maps.get(span["font"])
        for code, gid, (x, y), _bbox in span["chars"]:
            out.append({"x": x, "y": y, "ch": chr(code), "odia": gmap is not None,
                        "bold": "Bold" in span["font"], "glyph": gmap.get(gid) if gmap else None})
    return out


def within(c, col):
    lo, hi = COLS[col]
    return lo <= c["x"] < hi


def latin(chars, col):
    return "".join(c["ch"] for c in sorted((c for c in chars if not c["odia"] and within(c, col)),
                                           key=lambda c: c["x"])).strip()


def part_number(doc):
    """The roll's part number, printed top-left beside ଭାଗ."""
    digits = [c for c in page_chars(doc[0], {}) if c["y"] < 35 and 45 <= c["x"] < 75 and c["ch"].isdigit()]
    return "".join(c["ch"] for c in sorted(digits, key=lambda c: c["x"])) or None


def extract(pdf, decoder, pages=None):
    doc = pymupdf.open(pdf)
    font_maps = decoder.fonts(doc, pdf)
    part = part_number(doc)
    rows = {}
    for pno in range(doc.page_count):
        if pages and (pno + 1) not in pages:
            continue
        chars = page_chars(doc[pno], font_maps)
        serial_at = defaultdict(list)
        for c in chars:
            if not c["odia"] and c["y"] > 60 and within(c, "serial") and c["ch"].isdigit():
                serial_at[round(c["y"], 1)].append(c)
        for y, digits in serial_at.items():
            serial = int("".join(c["ch"] for c in sorted(digits, key=lambda c: c["x"])))
            band = [c for c in chars if y + BAND[0] <= c["y"] <= y + BAND[1]]
            odia = [c for c in band if c["odia"] and not c["bold"]]   # stream order = glyph order
            age = latin(band, "age")
            rows[serial] = {
                "serial": serial,
                "section": latin(band, "section"),
                "house": latin(band, "house"),
                "name": decoder.text([c for c in odia if within(c, "name")]),
                "reltype": latin(band, "reltype"),
                "relname": decoder.text([c for c in odia if within(c, "relname")]),
                "gender": latin(band, "gender"),
                "age": int(age) if age.isdigit() else None,
                "epic": latin(band, "epic").replace(" ", ""),
                "page": pno + 1,
            }
    return part, [rows[k] for k in sorted(rows)]


# ---------------------------------------------------------------- output

def write_sheet(path, part, rows, ward):
    wb = Workbook()
    ws = wb.active
    ws.title = "Voter List"
    ws.append(HEADERS)
    for cell in ws[1]:
        cell.font = Font(name="Calibri", bold=True, color="FFFFFFFF")
        cell.fill = PatternFill("solid", fgColor="FF2F5597")
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for r in rows:
        ws.append([
            r["house"],
            r["name"],
            r["reltype"],
            r["relname"],
            r["age"],
            r["gender"],
            f"{part}/{r['serial']}" if part else str(r["serial"]),
            r["epic"] or None,
            ward,
        ])
    for i, w in enumerate(WIDTHS):
        ws.column_dimensions[chr(65 + i)].width = w
    ws.freeze_panes = "A2"
    wb.save(path)


def dump_unknown(decoder, out_dir):
    """Render each unknown glyph to a PNG named by its hash, for labelling."""
    from fontTools.pens.basePen import BasePen
    from PIL import Image, ImageChops, ImageDraw

    class Outline(BasePen):
        def __init__(self, gs):
            super().__init__(gs)
            self.polys, self.cur = [], []

        def _moveTo(self, p):
            self.cur = [p]

        def _lineTo(self, p):
            self.cur.append(p)

        def _qCurveToOne(self, a, b):
            p0 = self.cur[-1]
            self.cur += [tuple((1 - t) ** 2 * p0[k] + 2 * (1 - t) * t * a[k] + t * t * b[k] for k in (0, 1))
                         for t in (i / 8 for i in range(1, 9))]

        def _curveToOne(self, a, b, c):
            p0 = self.cur[-1]
            self.cur += [tuple((1 - t) ** 3 * p0[k] + 3 * (1 - t) ** 2 * t * a[k] + 3 * (1 - t) * t * t * b[k]
                               + t ** 3 * c[k] for k in (0, 1)) for t in (i / 8 for i in range(1, 9))]

        def _closePath(self):
            self.polys.append(self.cur)
            self.cur = []

        _endPath = _closePath

    out = Path(out_dir)
    out.mkdir(parents=True, exist_ok=True)
    for h, (pdf, xref, gname) in decoder.unknown_seen.items():
        doc = pymupdf.open(pdf)
        gs = TTFont(io.BytesIO(doc.extract_font(xref)[3])).getGlyphSet()
        pen = Outline(gs)
        gs[gname].draw(pen)
        mask = Image.new("1", (300, 400), 0)
        for poly in pen.polys:
            if len(poly) > 2:
                layer = Image.new("1", mask.size, 0)
                ImageDraw.Draw(layer).polygon([(100 + x * 0.1, 280 - y * 0.1) for x, y in poly], fill=1)
                mask = ImageChops.logical_xor(mask, layer)
        img = Image.new("L", mask.size, 255)
        img.paste(0, (0, 0), mask)
        ImageDraw.Draw(img).line([(0, 280), (300, 280)], fill=180)
        img.save(out / f"{h}.png")
    print(f"  rendered {len(decoder.unknown_seen)} unknown glyph(s) to {out}")


def page_set(spec):
    if not spec:
        return None
    pages = set()
    for part in spec.split(","):
        a, _, b = part.partition("-")
        pages.update(range(int(a), int(b or a) + 1))
    return pages


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("pdf", nargs="+", help="electoral roll PDF(s); wildcards allowed")
    ap.add_argument("-o", "--out", help="output .xlsx (single input only; default: next to the PDF)")
    ap.add_argument("--ward", type=int, help="value for the Ward_No column (left blank if omitted)")
    ap.add_argument("--pages", help="pages to read, e.g. 3-10 or 1,3,5-7 (default: all)")
    ap.add_argument("--glyphs", default=HERE / "odia_glyphs.json", help="glyph table (default: alongside this script)")
    ap.add_argument("--dump-unknown", metavar="DIR", help="render glyphs missing from the table as PNGs")
    args = ap.parse_args()

    pdfs = [p for pattern in args.pdf for p in (glob.glob(pattern) or [pattern])]
    if args.out and len(pdfs) > 1:
        ap.error("-o can only be used with a single PDF")

    decoder = Decoder(args.glyphs)
    for pdf in pdfs:
        part, rows = extract(pdf, decoder, page_set(args.pages))
        out = args.out or str(Path(pdf).with_suffix(".xlsx"))
        write_sheet(out, part, rows, args.ward)
        gaps = [r["serial"] for r in rows if not r["name"] or r["age"] is None]
        print(f"{pdf}: part {part}, {len(rows)} voters -> {out}")
        if gaps:
            print(f"  check serials with a missing name or age: {gaps[:20]}")
    if decoder.unknown:
        total = sum(decoder.unknown.values())
        print(f"warning: {len(decoder.unknown)} glyph shape(s) not in the table ({total} occurrences); "
              f"their letters may lack vowel signs. Use --dump-unknown to label them.", file=sys.stderr)
        if args.dump_unknown:
            dump_unknown(decoder, args.dump_unknown)


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    main()
