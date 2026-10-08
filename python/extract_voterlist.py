"""
Extract an Odisha electoral-roll PDF into the K-Pulse voter upload sheet:

    house_no | name | relation | Relation_Name | Age | Gender | Assembly Part | Epic_no | Ward_No

Both printed layouts are read, and which one a PDF uses is detected from the
page: the tabular roll (one row per voter, as in P030028.pdf) and the photo-card
roll (three boxed cards a row, each line labelled ନାମ / ପିତାଙ୍କ ନାମ / ଘର ନଂ /
ବୟସ / ଲିଙ୍ଗ). A roll that has been rasterised, and so carries no text at all, is
read by OCR instead -- see ocr_cards.py, and treat its names as needing a check.

Why not just read the PDF text? The Odia names are set in an embedded
OROT-Mukta subset whose text layer maps every glyph back to a bare consonant:
vowel signs and conjuncts are lost ("ସାହୁ" extracts as "ସନହହ"). So this reads
the glyphs themselves. Each glyph is identified by a hash of its outline,
which stays the same across PDFs even though the glyph numbering changes, and
is looked up in odia_glyphs.json. The glyphs are then put back into logical
Unicode order: the pre-base ୋ/ୌ/ୈ, the reph and the conjunct parts.
Rolls that embed a font with a sound Unicode cmap skip all of that and are read
straight from the text layer.

Usage:
    pip install pymupdf fonttools openpyxl
    python extract_voterlist.py P030028.pdf                    -> P030028.xlsx
    python extract_voterlist.py P030028.pdf -o out.xlsx --ward 4 --pages 3-10
    python extract_voterlist.py roll.pdf --from 4 --to 6        (a page range)
    python extract_voterlist.py roll.pdf --from 4               (page 4 to the end)
    python extract_voterlist.py *.pdf                           (one sheet per PDF)

Glyphs missing from the table fall back to the PDF's own (lossy) letter and are
reported. Run with --dump-unknown DIR to render them as a PNG so they can be
added to odia_glyphs.json.
"""
import argparse
import glob
import hashlib
import io
import json
import re
import sys
from collections import Counter, defaultdict
from pathlib import Path

import pymupdf
from fontTools.ttLib import TTFont
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

HERE = Path(__file__).resolve().parent
HEADERS = ["house_no", "name", "relation", "Relation_Name", "Age", "Gender", "Assembly Part", "Epic_no", "Ward_No"]
WIDTHS = [16.5, 24, 9, 24, 8, 8, 16, 20, 9]

# Column x-ranges (PDF points) of the roll's table, and a row's vertical band
# relative to its serial number. Taken from the Kendrapara rolls.
COLS = {
    "serial": (25, 62), "section": (70, 110), "house": (120, 160),
    "name": (160, 300), "reltype": (270, 300), "relname": (305, 410),
    "gender": (410, 440), "age": (440, 470), "epic": (470, 575),
}
BAND = (-4, 14)

# ---------------------------------------------------------------- glyph → text

CONS = set("କଖଗଘଙଚଛଜଝଞଟଠଡଢଣତଥଦଧନପଫବଭମଯରଲଳଵଶଷସହୟୱଡ଼ଢ଼")
VOWELS = set("ଅଆଇଈଉଊଋଏଐଓଔ")
HALANT, E_SIGN, REPH = "୍", "େ", "ର୍"
# େ followed by a length mark is one two-part vowel sign.
TWO_PART = {"ା": "ୋ", "ୗ": "ୌ", "ୖ": "ୈ"}


ODIA_BLOCK = (0x0B00, 0x0B7F)


def is_odia(ch):
    return ODIA_BLOCK[0] <= ord(ch) <= ODIA_BLOCK[1]


def is_sign(ch):
    return "଼" <= ch <= "ୣ" or ch in "ଁଂଃ"


def to_logical(glyphs):
    """Glyph labels in visual order -> logical Unicode.

    Each item is (text, is_reph). The shaper draws the େ before its
    consonant, the reph after the base, and may draw ି before a subjoined
    consonant, so everything is regrouped per syllable cluster.
    """
    out, cl, pending_e = [], None, False

    def new():
        return {"cons": "", "marks": [], "reph": False, "e": False}

    def flush():
        nonlocal cl
        if cl is None:
            return
        text = (REPH if cl["reph"] else "") + cl["cons"]
        marks = cl["marks"]
        if cl["e"]:
            if marks and marks[0] in TWO_PART:
                text += TWO_PART[marks[0]]
                marks = marks[1:]
            else:
                text += E_SIGN
        out.append(text + "".join(marks))
        cl = None

    for text, reph in glyphs:
        if reph:
            if cl is None:
                cl = new()
            cl["reph"] = True
            if text:
                cl["marks"].append(text)
            continue
        if not text:
            continue
        if text == E_SIGN:
            pending_e = True
            continue
        first = text[0]
        if first in CONS:
            if cl is not None and cl["cons"].endswith(HALANT) and not cl["marks"]:
                cl["cons"] += text
            else:
                flush()
                cl = new()
                cl["cons"] = text
            if pending_e:
                cl["e"], pending_e = True, False
        elif first == HALANT:                      # subjoined consonant or bare halant
            if cl is None:
                cl = new()
            cl["cons"] += text
        elif is_sign(first):
            if cl is None:
                out.append(text)
            else:
                cl["marks"].append(text)
        else:
            flush()
            if first in VOWELS:
                cl = new()
                cl["cons"] = text
            else:
                out.append(text)
    flush()
    if pending_e:
        out.append(E_SIGN)
    return re.sub(r"\s+", " ", "".join(out)).strip()


def outline_hash(glyf, name):
    coords, ends, flags = glyf[name].getCoordinates(glyf)
    return hashlib.sha1(repr((list(coords), list(ends), [f & 1 for f in flags])).encode()).hexdigest()[:12]


class Decoder:
    def __init__(self, table_path):
        self.table = json.loads(Path(table_path).read_text(encoding="utf-8"))
        self.unknown = Counter()        # hash -> count
        self.unknown_seen = {}          # hash -> (pdf, font xref, glyph name, example)

    def fonts(self, doc, pdf):
        """font name -> {glyph id -> outline hash}, for every Odia font in the PDF."""
        maps = {}
        for page in doc:
            for xref, _ext, _type, base, *_ in page.get_fonts(full=True):
                name = base.split("+")[-1]
                if "OROT" not in name or name in maps:
                    continue
                font = TTFont(io.BytesIO(doc.extract_font(xref)[3]))
                glyf = font["glyf"]
                order = font.getGlyphOrder()
                maps[name] = {gid: (outline_hash(glyf, gn), (pdf, xref, gn)) for gid, gn in enumerate(order)}
        return maps

    def text(self, chars):
        if chars and all(c["glyph"] is None for c in chars):
            return re.sub(r"\s+", " ", "".join(c["ch"] for c in chars)).strip()
        glyphs = []
        for c in chars:
            if c["glyph"] is None:
                # Not one of the lossy subset fonts: its own text layer is sound.
                glyphs.append((c["ch"], False))
                continue
            h, origin = c["glyph"]
            label = self.table.get(h)
            if label is None:
                self.unknown[h] += 1
                self.unknown_seen.setdefault(h, origin)
                label = c["ch"]                    # the PDF's own guess: right consonant, no signs
            if label.startswith("@"):
                glyphs.append((label[1:], True))
            else:
                glyphs.append((label, False))
        return to_logical(glyphs)


# --------------------------------------------------------------- page → text

def page_chars(page, font_maps):
    out = []
    for span in page.get_texttrace():
        gmap = font_maps.get(span["font"])
        for code, gid, (x, y), _bbox in span["chars"]:
            out.append({"x": x, "y": y, "ch": chr(code), "odia": gmap is not None or is_odia(chr(code)),
                        "bold": "Bold" in span["font"], "glyph": gmap.get(gid) if gmap else None})
    return out


# ------------------------------------------------------------- page geometry

# A card's print lines are about 12pt apart and its labels sit a few points
# from their value, so these two tolerances separate lines from each other and
# a label from what follows it.
LINE_TOL = 5.0
CARD_GAP = 6.0


def lines_of(chars, tol=LINE_TOL):
    """Characters grouped into print lines, each sorted left to right."""
    out = []
    for c in sorted(chars, key=lambda c: c["y"]):
        if out and abs(c["y"] - out[-1][0]["y"]) <= tol:
            out[-1].append(c)
        else:
            out.append([c])
    return [sorted(line, key=lambda c: c["x"]) for line in out]


def runs(chars, gap=CARD_GAP):
    """A line split wherever a horizontal gap opens up, e.g. label | value."""
    out = []
    for c in chars:
        if out and c["x"] - out[-1][-1]["x"] <= gap:
            out[-1].append(c)
        else:
            out.append([c])
    return out


def first_letter(text):
    """The first real letter of a decoded label, used to tell labels apart.

    Only the leading consonant is needed, and that is the one part of a word
    the PDF's own text layer never loses, so a roll whose glyph shapes are
    missing from the table still files every value under the right field.
    """
    for ch in text:
        if ch in CONS or ch in VOWELS:
            return ch
    return ""


# ---------------------------------------------------------------- card layout

# Rolls printed as photo cards, three to a row, instead of one table row per
# voter. Nothing below is tied to one roll's measurements: the card is whatever
# box the page repeats, and each line is recognised by the first letter of its
# Odia label rather than by where it sits.
RELATION_BY_LABEL = {"ପ": "F", "ସ": "H", "ମ": "M"}   # ପିତାଙ୍କ / ସ୍ୱାମୀଙ୍କ / ମାତାଙ୍କ ନାମ
GENDER_BY_WORD = {"ପ": "M", "ସ": "F", "ଅ": "O"}      # ପୁରୁଷ / ସ୍ତ୍ରୀ / ଅନ୍ୟ
NAME_LABEL = "ନ"      # ନାମ
HOUSE_LABEL = "ଘ"     # ଘର ନଂ
AGE_LABEL = "ବ"       # ବୟସ
GENDER_LABEL = "ଲ"    # ଲିଙ୍ଗ, when a roll puts it on its own line
PHOTO_LABEL = "ଫ"     # ଫଟୋ ଉପଲବ୍ଧ, the caption printed in the photo box


def page_rects(page):
    """(cards, every box) on a page.

    The card border is the box size the page repeats most: cards are drawn
    identically thirty-odd times, while the page border occurs once and the
    photo box is far narrower. Finding the card this way means a roll with
    different margins, or a different number of cards per row, needs no change
    here.
    """
    rects, sizes, seen = [], Counter(), set()
    for drawing in page.get_drawings():
        r = drawing["rect"]
        key = (round(r.x0, 1), round(r.y0, 1), round(r.x1, 1), round(r.y1, 1))
        if key in seen or r.width < 20 or r.height < 20:
            continue
        seen.add(key)
        rects.append(r)
        if r.width >= 100 and r.height >= 40:
            sizes[(round(r.width), round(r.height))] += 1
    if not sizes:
        return [], rects
    (w, h), count = sizes.most_common(1)[0]
    if count < 3:                       # too few boxes to be a grid of cards
        return [], rects
    cards = [r for r in rects if abs(r.width - w) <= 2 and abs(r.height - h) <= 2]
    cards.sort(key=lambda r: (round(r.y0 / 10), r.x0))      # reading order
    return cards, rects


def photo_rect(card, rects):
    """The photo box inside a card, whose caption must not be read as a name."""
    nested = [r for r in rects
              if r.width > 20 and r.height > 20 and r.width < card.width and r.height < card.height
              and card.x0 <= r.x0 and r.x1 <= card.x1 and card.y0 <= r.y0 and r.y1 <= card.y1]
    return max(nested, key=lambda r: r.width * r.height, default=None)


def assign_to_cards(chars, cards, pad=2.0):
    """Each character to the single card it belongs to.

    The serial number is printed in a little box straddling the card's top
    edge, so the match has to be tolerant; picking the nearest centre keeps a
    glyph sitting on a shared border out of two cards at once.
    """
    buckets = [[] for _ in cards]
    for c in chars:
        best, best_distance = None, None
        for i, r in enumerate(cards):
            if not (r.x0 - pad <= c["x"] <= r.x1 + pad and r.y0 - pad <= c["y"] <= r.y1 + pad):
                continue
            distance = abs(c["x"] - (r.x0 + r.x1) / 2) + abs(c["y"] - (r.y0 + r.y1) / 2)
            if best is None or distance < best_distance:
                best, best_distance = i, distance
        if best is not None:
            buckets[best].append(c)
    return buckets


def after_colon(line, index=0):
    """The characters between the index-th colon on a line and the next one.

    "ବୟସ : 40 ଲିଙ୍ଗ : ସ୍ତ୍ରୀ" is one printed line carrying two fields, so each
    value is taken only as far as the following label. Where a roll omits the
    colons, the value is whatever follows the first gap instead.
    """
    colons = [i for i, c in enumerate(line) if c["ch"] == ":"]
    if index < len(colons):
        start = colons[index] + 1
        end = colons[index + 1] if index + 1 < len(colons) else len(line)
        return line[start:end]
    if index > 0:
        return []
    groups = runs(line)
    return [c for group in groups[1:] for c in group]


def value(line, decoder, index=0):
    """A field's characters, with any photo caption sharing the line dropped.

    The caption is excluded by the photo box as well, but not every roll draws
    one, so it is also recognised by its own label.
    """
    out = []
    for group in runs(after_colon(line, index)):
        odia = [c for c in group if c["odia"]]
        if odia and first_letter(decoder.text(odia)) == PHOTO_LABEL:
            break
        out.extend(group)
    return out


def odia_text(chars, decoder):
    return decoder.text([c for c in chars if c["odia"] or c["ch"].isspace()])


def latin_text(chars):
    return "".join(c["ch"] for c in chars if not c["odia"]).strip()


def parse_card(card, inside, rects, decoder):
    """One card -> one voter, read by label rather than by position."""
    photo = photo_rect(card, rects)
    if photo is not None:
        inside = [c for c in inside
                  if not (photo.x0 - 1 <= c["x"] <= photo.x1 + 1
                          and photo.y0 <= c["y"] <= photo.y1 + 1)]

    row = {"serial": None, "section": "", "house": "", "name": "", "reltype": "",
           "relname": "", "gender": "", "age": None, "epic": ""}
    for line in lines_of(inside):
        odia = [c for c in line if c["odia"]]
        if not odia:
            # The card's top line: serial in its own box on the left, EPIC right.
            groups = runs(line)
            serial = latin_text(groups[0])
            if serial.isdigit():
                row["serial"] = int(serial)
            if len(groups) > 1:
                row["epic"] = latin_text(groups[-1]).replace(" ", "")
            continue

        label = first_letter(decoder.text(odia))
        if label == PHOTO_LABEL:
            continue
        if label == NAME_LABEL:
            row["name"] = odia_text(value(line, decoder), decoder)
        elif label in RELATION_BY_LABEL:
            row["reltype"] = RELATION_BY_LABEL[label]
            row["relname"] = odia_text(value(line, decoder), decoder)
        elif label == HOUSE_LABEL:
            row["house"] = latin_text(value(line, decoder))
        elif label == AGE_LABEL:
            age = "".join(c["ch"] for c in value(line, decoder) if c["ch"].isdigit())
            row["age"] = int(age) if age else None
            # Age and gender usually share a line, where gender follows the
            # second colon; rolls that split them are handled by ଲିଙ୍ଗ below.
            gender = odia_text(value(line, decoder, 1), decoder)
            if gender:
                row["gender"] = GENDER_BY_WORD.get(first_letter(gender), "")
        elif label == GENDER_LABEL:
            row["gender"] = GENDER_BY_WORD.get(
                first_letter(odia_text(value(line, decoder), decoder)), "")
    return row


def extract_cards(doc, decoder, font_maps, pages):
    rows = []
    for pno in range(doc.page_count):
        if pages and (pno + 1) not in pages:
            continue
        page = doc[pno]
        cards, rects = page_rects(page)
        if not cards:
            continue
        chars = page_chars(page, font_maps)
        for card, inside in zip(cards, assign_to_cards(chars, cards)):
            row = parse_card(card, inside, rects, decoder)
            if row["name"] or row["serial"] is not None:   # blank cards pad the last page
                row["page"] = pno + 1
                rows.append(row)
    return rows


# --------------------------------------------------------------- table layout

def within(c, col):
    lo, hi = COLS[col]
    return lo <= c["x"] < hi


def latin(chars, col):
    return "".join(c["ch"] for c in sorted((c for c in chars if not c["odia"] and within(c, col)),
                                           key=lambda c: c["x"])).strip()


def extract_table(doc, decoder, font_maps, pages):
    rows = {}
    for pno in range(doc.page_count):
        if pages and (pno + 1) not in pages:
            continue
        chars = page_chars(doc[pno], font_maps)
        serial_at = defaultdict(list)
        for c in chars:
            if not c["odia"] and c["y"] > 60 and within(c, "serial") and c["ch"].isdigit():
                serial_at[round(c["y"], 1)].append(c)
        for y, digits in serial_at.items():
            serial = int("".join(c["ch"] for c in sorted(digits, key=lambda c: c["x"])))
            band = [c for c in chars if y + BAND[0] <= c["y"] <= y + BAND[1]]
            odia = [c for c in band if c["odia"] and not c["bold"]]   # stream order = glyph order
            age = latin(band, "age")
            rows[serial] = {
                "serial": serial,
                "section": latin(band, "section"),
                "house": latin(band, "house"),
                "name": decoder.text([c for c in odia if within(c, "name")]),
                "reltype": latin(band, "reltype"),
                "relname": decoder.text([c for c in odia if within(c, "relname")]),
                "gender": latin(band, "gender"),
                "age": int(age) if age.isdigit() else None,
                "epic": latin(band, "epic").replace(" ", ""),
                "page": pno + 1,
            }
    return [rows[k] for k in sorted(rows)]


# -------------------------------------------------------------- scanned roll

# "ସ୍ୱାମୀଙ୍କ ନାମ: X" sometimes comes back with the colon a word early, which
# leaves ନାମ at the front of the value. Only ନାମ as a word of its own is dropped,
# so a name that merely begins that way, ନାମିତା for one, is left alone.
LABEL_LEFTOVER = re.compile(r"^ନାମ[଀-୿]?\s+")


def clean_odia(text):
    """A name as OCR returned it, with stray marks it invented dropped."""
    kept = "".join(ch for ch in text if is_odia(ch) or ch.isspace())
    return LABEL_LEFTOVER.sub("", re.sub(r"\s+", " ", kept).strip())


def card_row_from_text(serial, epic, text, ocr):
    """One OCR'd card -> one voter, read by label exactly as a digital roll is."""
    row = {"serial": None, "section": "", "house": "", "name": "", "reltype": "",
           "relname": "", "gender": "", "age": None, "epic": epic.strip()}
    digits = "".join(ch for ch in serial if ch.isdigit())
    if digits:
        row["serial"] = int(digits)

    for line in text.splitlines():
        line = line.strip()
        if not line:
            continue
        label = ocr.first_letter(line)
        if not label or label == PHOTO_LABEL:
            continue
        if label == NAME_LABEL:
            row["name"] = clean_odia(ocr.value_after(line))
        elif label in RELATION_BY_LABEL:
            row["reltype"] = RELATION_BY_LABEL[label]
            row["relname"] = clean_odia(ocr.value_after(line))
        elif label == HOUSE_LABEL:
            found = ocr.DIGITS.search(ocr.value_after(line))
            row["house"] = found.group(0) if found else ""
        elif label == AGE_LABEL:
            found = ocr.DIGITS.search(ocr.value_after(line))
            row["age"] = int(found.group(0)) if found else None
            row["gender"] = gender_on(line, ocr)
        elif label == GENDER_LABEL and not row["gender"]:
            row["gender"] = gender_on(line, ocr)
    return row


def gender_on(line, ocr):
    """ପୁରୁଷ or ସ୍ତ୍ରୀ at the end of a line, by its first letter.

    Only that letter is needed, which matters here: OCR often clips the rest of
    so short a word, and ଲିଙ୍ଗ's own colon is the one most often lost.
    """
    for candidate in (ocr.value_after(line, 1), ocr.last_odia_run(line)):
        gender = GENDER_BY_WORD.get(ocr.first_letter(candidate), "")
        if gender:
            return gender
    return ""


def extract_scanned(doc, pages, lang, tesseract):
    """A roll whose pages are images: every field has to be read by OCR."""
    import ocr_cards as ocr

    ocr.use_tesseract(tesseract)
    rows, part = [], None
    for pno in range(doc.page_count):
        if pages and (pno + 1) not in pages:
            continue
        gray, boxes, cards = ocr.prepared(doc[pno])
        # OCR takes the better part of a minute a page and prints nothing until
        # the whole roll is done, which looks like a hang. So say where we are.
        print(f"  page {pno + 1} of {doc.page_count}: {len(cards)} card(s)",
              file=sys.stderr, flush=True)
        if not cards:
            continue
        if part is None:
            header = ocr.read(gray, (0, 0, gray.shape[1], cards[0][1]), ocr.BLOCK, lang)
            part = ocr.part_number(header)
        for card in cards:
            serial, epic, text = ocr.card_fields(gray, card, boxes, lang)
            row = card_row_from_text(serial, epic, text, ocr)
            if row["name"] or row["serial"] is not None:
                row["page"] = pno + 1
                row["ocr"] = True
                rows.append(row)
    return part, rows


def out_of_sequence(rows):
    """Serials that do not follow the one before them.

    A roll numbers its cards in order, so a break is a misread rather than a
    gap in the roll -- which makes this the one check that catches an OCR'd
    serial being wrong without anything to compare against.
    """
    odd, previous = [], None
    for row in rows:
        serial = row["serial"]
        if serial is None or (previous is not None and serial != previous + 1):
            odd.append(serial)
        if serial is not None:
            previous = serial
    return odd


# ------------------------------------------------------------------- dispatch

PART_LABEL = "ଭ"      # ଭାଗ, the part number in the page header
PART_REACH = 80.0     # how far right of that label the number can be printed


def part_number(doc, decoder, font_maps):
    """The roll's part number, printed beside ଭାଗ in the header of page one.

    Only the digits just to the right of that label count: the same header also
    carries the constituency and polling station numbers, which would otherwise
    be read as part of it.
    """
    head = [c for c in page_chars(doc[0], font_maps) if c["y"] < 60]
    for line in lines_of(head):
        groups = runs(line)
        for i, group in enumerate(groups):
            odia = [c for c in group if c["odia"]]
            if not odia or first_letter(decoder.text(odia)) != PART_LABEL:
                continue
            limit = group[-1]["x"] + PART_REACH
            digits = ""
            for c in [c for c in group if not c["odia"]] + [c for g in groups[i + 1:] for c in g]:
                if c["x"] > limit:
                    break
                if c["ch"].isdigit():
                    digits += c["ch"]
                elif digits:                   # the number has ended
                    break
            if digits:
                return digits
    return None


def looks_like_cards(doc, pages):
    """Whether this roll prints a photo card per voter rather than a table row."""
    for pno in range(doc.page_count):
        if pages and (pno + 1) not in pages:
            continue
        cards, _rects = page_rects(doc[pno])
        if cards:
            return True
    return False


def extract(pdf, decoder, pages=None, lang="ori", tesseract=None):
    """Voters from any of the roll layouts; which one it is, is read off the page.

    A roll with no text layer at all has been rasterised somewhere along the
    way and can only be read by OCR, which is a good deal less certain than the
    other two paths -- see the warning extract_voterlist prints.
    """
    doc = pymupdf.open(pdf)
    if pages and not any(page + 1 in pages for page in range(doc.page_count)):
        print(f"warning: {pdf} has {doc.page_count} page(s), which is outside the pages asked for",
              file=sys.stderr)
    if all(not page.get_text().strip() for page in doc):
        return extract_scanned(doc, pages, lang, tesseract)
    font_maps = decoder.fonts(doc, pdf)
    part = part_number(doc, decoder, font_maps)
    if looks_like_cards(doc, pages):
        return part, extract_cards(doc, decoder, font_maps, pages)
    return part, extract_table(doc, decoder, font_maps, pages)


# ---------------------------------------------------------------- output

def write_sheet(path, part, rows, ward):
    wb = Workbook()
    ws = wb.active
    ws.title = "Voter List"
    ws.append(HEADERS)
    for cell in ws[1]:
        cell.font = Font(name="Calibri", bold=True, color="FFFFFFFF")
        cell.fill = PatternFill("solid", fgColor="FF2F5597")
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for r in rows:
        ws.append([
            r["house"],
            r["name"],
            r["reltype"],
            r["relname"],
            r["age"],
            r["gender"],
            f"{part}/{r['serial']}" if part else str(r["serial"]),
            r["epic"] or None,
            ward,
        ])
    for i, w in enumerate(WIDTHS):
        ws.column_dimensions[chr(65 + i)].width = w
    ws.freeze_panes = "A2"
    wb.save(path)


def dump_unknown(decoder, out_dir):
    """Render each unknown glyph to a PNG named by its hash, for labelling."""
    from fontTools.pens.basePen import BasePen
    from PIL import Image, ImageChops, ImageDraw

    class Outline(BasePen):
        def __init__(self, gs):
            super().__init__(gs)
            self.polys, self.cur = [], []

        def _moveTo(self, p):
            self.cur = [p]

        def _lineTo(self, p):
            self.cur.append(p)

        def _qCurveToOne(self, a, b):
            p0 = self.cur[-1]
            self.cur += [tuple((1 - t) ** 2 * p0[k] + 2 * (1 - t) * t * a[k] + t * t * b[k] for k in (0, 1))
                         for t in (i / 8 for i in range(1, 9))]

        def _curveToOne(self, a, b, c):
            p0 = self.cur[-1]
            self.cur += [tuple((1 - t) ** 3 * p0[k] + 3 * (1 - t) ** 2 * t * a[k] + 3 * (1 - t) * t * t * b[k]
                               + t ** 3 * c[k] for k in (0, 1)) for t in (i / 8 for i in range(1, 9))]

        def _closePath(self):
            self.polys.append(self.cur)
            self.cur = []

        _endPath = _closePath

    out = Path(out_dir)
    out.mkdir(parents=True, exist_ok=True)
    for h, (pdf, xref, gname) in decoder.unknown_seen.items():
        doc = pymupdf.open(pdf)
        gs = TTFont(io.BytesIO(doc.extract_font(xref)[3])).getGlyphSet()
        pen = Outline(gs)
        gs[gname].draw(pen)
        mask = Image.new("1", (300, 400), 0)
        for poly in pen.polys:
            if len(poly) > 2:
                layer = Image.new("1", mask.size, 0)
                ImageDraw.Draw(layer).polygon([(100 + x * 0.1, 280 - y * 0.1) for x, y in poly], fill=1)
                mask = ImageChops.logical_xor(mask, layer)
        img = Image.new("L", mask.size, 255)
        img.paste(0, (0, 0), mask)
        ImageDraw.Draw(img).line([(0, 280), (300, 280)], fill=180)
        img.save(out / f"{h}.png")
    print(f"  rendered {len(decoder.unknown_seen)} unknown glyph(s) to {out}")


def page_set(spec):
    if not spec:
        return None
    pages = set()
    for part in spec.split(","):
        a, _, b = part.partition("-")
        pages.update(range(int(a), int(b or a) + 1))
    return pages


# --to may be left off to mean "to the end of the roll", and the page count is
# not known until the PDF is open, so the range simply runs past any of them.
OPEN_ENDED = 1_000_000


def page_range(first, last):
    """The pages --from and --to select, with either end allowed to be open."""
    if first is None and last is None:
        return None
    return range(first or 1, (last or OPEN_ENDED) + 1)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("pdf", nargs="+", help="electoral roll PDF(s); wildcards allowed")
    ap.add_argument("-o", "--out", help="output .xlsx (single input only; default: next to the PDF)")
    ap.add_argument("--ward", type=int, help="value for the Ward_No column (left blank if omitted)")
    ap.add_argument("--pages", help="pages to read, e.g. 3-10 or 1,3,5-7 (default: all)")
    ap.add_argument("--from", dest="from_page", type=int, metavar="N",
                    help="first page to read, counting from 1 (default: the first)")
    ap.add_argument("--to", dest="to_page", type=int, metavar="N",
                    help="last page to read, inclusive (default: the last)")
    ap.add_argument("--glyphs", default=HERE / "odia_glyphs.json", help="glyph table (default: alongside this script)")
    ap.add_argument("--dump-unknown", metavar="DIR", help="render glyphs missing from the table as PNGs")
    ap.add_argument("--lang", default="ori", help="tesseract language for scanned rolls (default: ori)")
    ap.add_argument("--tesseract", help="path to tesseract.exe, if it is not on PATH")
    args = ap.parse_args()

    pdfs = [p for pattern in args.pdf for p in (glob.glob(pattern) or [pattern])]
    if args.out and len(pdfs) > 1:
        ap.error("-o can only be used with a single PDF")

    if args.pages and (args.from_page or args.to_page):
        ap.error("--pages and --from/--to do the same job; use one or the other")
    for name, value in (("--from", args.from_page), ("--to", args.to_page)):
        if value is not None and value < 1:
            ap.error(f"{name} counts from 1")
    if args.from_page and args.to_page and args.from_page > args.to_page:
        ap.error(f"--from {args.from_page} is after --to {args.to_page}")
    pages = page_set(args.pages) or page_range(args.from_page, args.to_page)

    missing = [pdf for pdf in pdfs if not Path(pdf).is_file()]
    if missing:
        ap.error("no such file: " + ", ".join(missing)
                 + "\n(the usage examples call it roll.pdf; use your own file's name)")

    decoder = Decoder(args.glyphs)
    failures = 0
    for pdf in pdfs:
        try:
            part, rows = extract(pdf, decoder, pages, args.lang, args.tesseract)
        except Exception as error:               # one bad PDF must not lose the rest
            failures += 1
            print(f"{pdf}: could not be read: {error}", file=sys.stderr)
            continue
        out = args.out or str(Path(pdf).with_suffix(".xlsx"))
        write_sheet(out, part, rows, args.ward)
        gaps = [r["serial"] for r in rows if not r["name"] or r["age"] is None]
        print(f"{pdf}: part {part}, {len(rows)} voters -> {out}")
        if gaps:
            print(f"  check serials with a missing name or age: {gaps[:20]}")
        if rows and rows[0].get("ocr"):
            broken = out_of_sequence(rows)
            if broken:
                print(f"  check these serials, which break the run: {broken[:20]}")
            print("  read by OCR: spot-check the names against the PDF before loading them")
    if decoder.unknown:
        total = sum(decoder.unknown.values())
        print(f"warning: {len(decoder.unknown)} glyph shape(s) not in the table ({total} occurrences); "
              f"their letters may lack vowel signs. Use --dump-unknown to label them.", file=sys.stderr)
        if args.dump_unknown:
            dump_unknown(decoder, args.dump_unknown)
    return 1 if failures else 0


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    sys.exit(main())
